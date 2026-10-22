"""数据导出服务（需求 §28/§29/§30/§31）。

- JSON：完整数据（渠道/账户/分类/交易/设置），金额存字符串避免精度问题；
- CSV：交易数据（中文表头，UTF-8 BOM 便于 Excel 直接打开）；
- XLSX：多 Sheet（Transactions / Channels / Accounts / Categories），
  金额与余额写数字（Excel 可计算），日期写文本（避免格式歧义）。
"""

from __future__ import annotations

import csv
import json
import logging
from datetime import date, datetime, timedelta
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import sessionmaker

from database import database
from database.models import Account, Category, Channel, Setting, Transaction
from database.repositories.transaction_repository import TransactionRepository

logger = logging.getLogger(__name__)

CSV_HEADERS = ["日期", "类型", "金额", "渠道", "账户", "分类", "备注", "标签"]


def _counts(channels: int, accounts: int, categories: int, transactions: int) -> dict:
    return {
        "channels": channels,
        "accounts": accounts,
        "categories": categories,
        "transactions": transactions,
    }


def month_export_range(
    year: int, month: int, today: date | None = None
) -> tuple[date, date]:
    """指定月份的导出范围（含首尾）。

    当月未过完时导出「1 日 ~ 今天」，其余月份导出整月。
    """
    today = today or date.today()
    first_day = date(year, month, 1)
    if (year, month) == (today.year, today.month):
        return first_day, today
    next_month = date(year + (month == 12), month % 12 + 1, 1)
    return first_day, next_month - timedelta(days=1)


def _validate_date_range(date_from: date | None, date_to: date | None) -> None:
    """起止日期（含首尾）；允许起止相同，但不允许开始晚于结束。"""
    if date_from is not None and date_to is not None and date_from > date_to:
        raise ValueError("开始日期不能晚于结束日期")


class ExportService:
    """数据导出：UI 只经由本服务写出文件。"""

    def __init__(self, session_factory: sessionmaker | None = None) -> None:
        self._session_factory = session_factory or database.get_session_factory()

    def export_json(self, path: Path) -> dict:
        """导出完整 JSON（需求 §29）。金额以字符串保存；返回各实体数量。"""
        with database.session_scope(self._session_factory) as session:
            transactions = TransactionRepository(session).list_all()
            channels = session.scalars(select(Channel).order_by(Channel.name)).all()
            accounts = session.scalars(select(Account).order_by(Account.name)).all()
            categories = session.scalars(
                select(Category).order_by(Category.type, Category.name)
            ).all()
            settings = session.scalars(select(Setting)).all()

            payload = {
                "app": "MoneyManager",
                "version": 1,
                "exported_at": datetime.now().isoformat(timespec="seconds"),
                "channels": [
                    {
                        "name": c.name,
                        "description": c.description,
                        "monthly_amount": str(c.monthly_amount),
                        "is_active": c.is_active,
                    }
                    for c in channels
                ],
                "accounts": [
                    {
                        "name": a.name,
                        "description": a.description,
                        "initial_balance": str(a.initial_balance),
                        "is_active": a.is_active,
                    }
                    for a in accounts
                ],
                "categories": [
                    {
                        "name": c.name,
                        "type": c.type,
                        "description": c.description,
                        "is_active": c.is_active,
                    }
                    for c in categories
                ],
                "transactions": [
                    {
                        "type": t.type,
                        "amount": str(t.amount),
                        "date": t.date.isoformat(),
                        "channel": t.channel.name if t.channel else "",
                        "account": t.account.name if t.account else "",
                        "category": t.category.name if t.category else "",
                        "note": t.note,
                        "tags": t.tags,
                    }
                    for t in transactions
                ],
                "settings": {s.key: s.value for s in settings},
            }

        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        logger.info("导出 JSON：%s（交易 %d 笔）", path, len(payload["transactions"]))
        return _counts(len(channels), len(accounts), len(categories), len(transactions))

    def export_csv(
        self, path: Path, *, date_from: date | None = None, date_to: date | None = None
    ) -> dict:
        """导出交易 CSV（需求 §30）。

        date_from / date_to 为可选的日期范围（含首尾）；均不传时导出全部交易。
        """
        _validate_date_range(date_from, date_to)
        with database.session_scope(self._session_factory) as session:
            transactions = self._query_transactions(session, date_from, date_to)
            rows = [
                [
                    t.date.isoformat(),
                    t.type,
                    str(t.amount),
                    t.channel.name if t.channel else "",
                    t.account.name if t.account else "",
                    t.category.name if t.category else "",
                    t.note,
                    t.tags,
                ]
                for t in transactions
            ]

        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(CSV_HEADERS)
            writer.writerows(rows)
        logger.info(
            "导出 CSV：%s（交易 %d 笔，范围 %s ~ %s）",
            path, len(rows),
            date_from.isoformat() if date_from else "不限",
            date_to.isoformat() if date_to else "不限",
        )
        return _counts(0, 0, 0, len(rows))

    def export_xlsx(
        self, path: Path, *, date_from: date | None = None, date_to: date | None = None
    ) -> dict:
        """导出多 Sheet Excel（需求 §31）。

        Transactions 按可选日期范围（含首尾）过滤；实体 Sheet 保持全量。
        """
        from openpyxl import Workbook

        _validate_date_range(date_from, date_to)
        with database.session_scope(self._session_factory) as session:
            transactions = self._query_transactions(session, date_from, date_to)
            channels = session.scalars(select(Channel).order_by(Channel.name)).all()
            accounts = session.scalars(select(Account).order_by(Account.name)).all()
            categories = session.scalars(
                select(Category).order_by(Category.type, Category.name)
            ).all()

            workbook = Workbook()

            ws_tx = workbook.active
            ws_tx.title = "Transactions"
            ws_tx.append(CSV_HEADERS)
            for t in transactions:
                ws_tx.append(
                    [
                        t.date.isoformat(),
                        t.type,
                        float(t.amount),  # 数字单元格：Excel 可计算；2 位小数金额可无损往返
                        t.channel.name if t.channel else "",
                        t.account.name if t.account else "",
                        t.category.name if t.category else "",
                        t.note,
                        t.tags,
                    ]
                )

            ws_ch = workbook.create_sheet("Channels")
            ws_ch.append(["名称", "描述", "每月固定金额", "启用"])
            for c in channels:
                ws_ch.append(
                    [c.name, c.description, float(c.monthly_amount), c.is_active]
                )

            ws_ac = workbook.create_sheet("Accounts")
            ws_ac.append(["名称", "描述", "初始余额", "启用"])
            for a in accounts:
                ws_ac.append([a.name, a.description, float(a.initial_balance), a.is_active])

            ws_cat = workbook.create_sheet("Categories")
            ws_cat.append(["名称", "类型", "描述", "启用"])
            for c in categories:
                ws_cat.append([c.name, c.type, c.description, c.is_active])

            path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(path)

        logger.info(
            "导出 XLSX：%s（交易 %d 笔，范围 %s ~ %s）",
            path, len(transactions),
            date_from.isoformat() if date_from else "不限",
            date_to.isoformat() if date_to else "不限",
        )
        return _counts(len(channels), len(accounts), len(categories), len(transactions))

    @staticmethod
    def _query_transactions(session, date_from: date | None, date_to: date | None):
        """按可选范围查询交易；都不传时返回全部。"""
        repository = TransactionRepository(session)
        if date_from is None and date_to is None:
            return repository.list_all()
        return repository.list_in_range(date_from, date_to)
